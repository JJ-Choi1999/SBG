from docx import Document
from docx.enum.dml import MSO_THEME_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.opc.oxml import qn


def extra_hyperlink(para: Paragraph, para_text: str, docx: Document):
    """
    提取超链接
    :param para: docx 段落对象
    :param para_text: docx 段落文本
    :param docx: docx 对象
    :return:
    """
    if not para._p.xpath('.//w:hyperlink'): return {}
    hl = Hyperlink(para._p.xpath('.//w:hyperlink')[0], docx)
    return {
        'hl_text': hl.text,
        'hl_url': hl.url,
        'hl_runs': hl.runs,
        'hl_index': para_text.find(hl.text)
    }

def extra_paragraph(docx_path):
    """
    提取docx段落
    :param docx_path:
    :return:
    """
    para_texts = []
    docx = Document(docx_path)

    for para in docx.paragraphs:

        hyperlink_result = extra_hyperlink(para, para.text, docx)
        hl_index = hyperlink_result.get('hl_index')
        hl_url = hyperlink_result.get('hl_url')
        is_flag = False

        # 假如存在超链接URL但是docx段落中无法找到, 表示是pdf2docx 转换时超链接非标准处理, 故迭代段落run补全超链接
        if hl_index == -1 and hl_url:
            para_text = ''
            for run in para.runs:
                if not is_flag and not run.text:
                    run.text = hl_url
                    is_flag = True
                para_text += run.text
        para_texts.append(para.text)

    docx.save(docx_path)

    return para_texts

def extra_table(docx_path):
    cell_texts = []
    docx = Document(docx_path)
    for table in docx.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_texts.append(cell.text)

    return cell_texts

if __name__ == '__main__':
    docx_path = r"C:\Users\Lenovo\Desktop\Clearstream FAQs – Clearing mandate for U.S. Treasury securities – U.S.A__1.docx"
    print(extra_paragraph(docx_path))