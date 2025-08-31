import os
import uuid
import markdown

from docx import Document
from docx.oxml.ns import qn
from bs4 import BeautifulSoup

def set_header(elem, docx):
    """
    设置标题
    :param elem:
    :param docx:
    :return:
    """
    header = None

    if elem.name == 'h1':
        header = docx.add_heading(elem.text, level=1)
    elif elem.name == 'h2':
        header = docx.add_heading(elem.text, level=2)
    elif elem.name == 'h3':
        header = docx.add_heading(elem.text, level=3)
    elif elem.name == 'h4':
        header = docx.add_heading(elem.text, level=4)
    elif elem.name == 'h5':
        header = docx.add_heading(elem.text, level=5)
    elif elem.name == 'h6':
        header = docx.add_heading(elem.text, level=6)

    return header

def set_paragraph(elem, docx):
    """
    设置段落
    :param elem:
    :param docx:
    :return:
    """
    paragraph = docx.add_paragraph()

    for child in elem.children:

        if child.name == 'strong':
            run = paragraph.add_run(child.text)
            run.bold = True

        elif child.name == 'em':
            run = paragraph.add_run(child.text)
            run.italic = True

        else:
            run = paragraph.add_run(child.text)

    return paragraph

def set_table(elem, docx):
    """
    设置表格
    :param elem:
    :param docx:
    :return:
    """
    # 处理表格需要对单元格样式进行进一步拓展
    table = docx.add_table(row=1, cols=len(elem.find('tr').find_all('th')))
    table.style = 'Table Grid'
    header_cells = table.rows[0].cell

    # 表头处理
    for i, th in enumerate(elem.find('tr').find_all('th')):
        header_cells[i].text = th.text

    # 表body处理
    for tr in elem.find_all('tr')[1:]:
        row_cells = table.add_row().cells
        for i, td in enumerate(tr.find_all('td')):
            row_cell = row_cells[i]
            row_cell.text = td.text

    return table

def md_to_docx(md_text: str, docx_path: str = None) -> str:
    if not docx_path:
        docx_path = os.path.join(os.getcwd(), f'{str(uuid.uuid1())}.docx')

    # 1. md -> html
    html_obj = markdown.markdown(
        md_text,
        extensions=[
            'markdown.extensions.extra',
            'markdown.extensions.tables',
            'markdown.extensions.fenced_code'
        ] # 支持额外的markdown -> html 语法转换
    )

    # 创建 docx 对象
    docx = Document()
    # 获取默认样式
    style = docx.styles['Normal']
    # 设置英文字体
    style.font.name = 'Arial'
    # 设置中文字体
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # html 解析
    soup = BeautifulSoup(html_obj, features='html.parser')

    for elem in soup.children:

        if elem.name in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
            header = set_header(elem, docx)

        elif elem.name == 'p':
            paragraph = set_paragraph(elem, docx)

        elif elem.name == 'ul':
            for li in elem.find_all('li'):
                ul_li = docx.add_paragraph(li.text, style='List Bullet')

        elif elem.name == 'ol':
            for idx, li in enumerate(elem.find_all('li')):
                ol_li = docx.add_paragraph(f'{idx + 1}. {li.text}', style='List Number')

        elif elem.name == 'table':
            table = set_table(elem, docx)

    return docx_path